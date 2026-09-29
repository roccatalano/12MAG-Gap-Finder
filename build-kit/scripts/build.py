import sys; sys.path.insert(0,'/tmp/w'); from data import *
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import ColorScaleRule, CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter as L

F="Arial"; NAVY="1F3A5F"
hdr_font=Font(name=F,bold=True,color="FFFFFF"); hdr_fill=PatternFill("solid",fgColor=NAVY)
inp_fill=PatternFill("solid",fgColor="FFF2CC"); base=Font(name=F,size=10)
thin=Side(style="thin",color="BFBFBF")
wb=Workbook()

HEAD0=["Paper ID","Year","Source","Exam","Question","Marks","Topic code","Topic","Area of study","Secondary code","Task type","Correct answer","Success rate","Difficulty","Response spread","Notes","Marks lost","Exam % lost","First row of paper"]
HEAD=HEAD0+["Textbook sections to redo (Cambridge)","Video solution (MaffsGuru)"]
WID=[16,7,13,8,9,7,10,48,26,11,9,9,10,10,30,48,10,10,9,40,16]
N=3000
def rows_for(pid,year,src,exam):
    out=[]
    if exam=="Exam 1":
        for q,ans,a,b,c,d,code,sec,task,note in E1:
            p=dict(A=a,B=b,C=c,D=d); wrong=max((k for k in p if k!=ans),key=p.get)
            out.append([pid,year,src,exam,str(q),1,code,None,None,sec or None,task,ans,p[ans]/100,None,
                        f"Most common wrong answer: {wrong} ({p[wrong]}%)",note])
    else:
        for q,m,dist,code,sec,task,note in E2:
            succ=round(sum(k*v for k,v in enumerate(dist))/100/m,3)
            out.append([pid,year,src,exam,q,m,code,None,None,sec or None,task,None,succ,None,
                        " | ".join(f"{k} marks: {v}%" for k,v in enumerate(dist)),note])
    return out

def formulas(r):
    return {8:f'=IFERROR(INDEX(Codes!$D:$D,MATCH(G{r},Codes!$A:$A,0)),"")',
            9:f'=IFERROR(INDEX(Codes!$C:$C,MATCH(G{r},Codes!$A:$A,0)),"")',
            14:f'=IF(M{r}="","",IF(M{r}>=0.7,"Easy",IF(M{r}>=0.4,"Medium","Hard")))',
            17:f'=IF(M{r}="","",F{r}*(1-M{r}))',
            18:f'=IF(M{r}="","",Q{r}/SUMIF($A:$A,A{r},$F:$F))',
            19:f'=IF(COUNTIF(A$2:A{r},A{r})=1,1,0)',
            20:f'=IFERROR(INDEX(Codes!$E:$E,MATCH(G{r},Codes!$A:$A,0)),"")'}

from video import link as VL
LINKS=[]
def data_sheet(ws,rows,table_name):
    for c,h in enumerate(HEAD,1):
        x=ws.cell(1,c,h); x.font=hdr_font; x.fill=hdr_fill; x.alignment=Alignment(wrap_text=True,vertical="center")
        ws.column_dimensions[L(c)].width=WID[c-1]
    for i,row in enumerate(rows):
        r=i+2
        for c,v in enumerate(row,1): ws.cell(r,c,v)
        for c,f in formulas(r).items(): ws.cell(r,c,f)
        u,lab=VL(row[0],row[4])
        if u:
            x=ws.cell(r,21,"▶ "+("Full video" if lab=="whole video" else lab)); x.hyperlink=u; LINKS.append((ws.title,r))
        for c in range(1,22):
            x=ws.cell(r,c); x.font=base
        for c in (1,2,3,4,5,6,7,10,11,12,13,15,16): ws.cell(r,c).fill=inp_fill
        ws.cell(r,21).font=Font(name=F,size=10,color="0563C1",underline="single")
        ws.cell(r,13).number_format="0%"; ws.cell(r,17).number_format="0.00"; ws.cell(r,18).number_format="0.0%"
    last=len(rows)+1
    t=Table(displayName=table_name,ref=f"A1:U{last}")
    t.tableStyleInfo=TableStyleInfo(name="TableStyleLight9",showRowStripes=True); ws.add_table(t)
    ws.freeze_panes="F2"; ws.row_dimensions[1].height=30
    ws.conditional_formatting.add(f"M2:M{N}",ColorScaleRule(start_type="num",start_value=0.2,start_color="F8696B",mid_type="num",mid_value=0.55,mid_color="FFEB84",end_type="num",end_value=0.9,end_color="63BE7B"))
    for col,lst in (("G","=Codes!$A$2:$A$60"),("J","=Codes!$A$2:$A$60"),("K","=Lists!$A$2:$A$5"),("C","=Lists!$B$2:$B$6"),("D","=Lists!$C$2:$C$3")):
        dv=DataValidation(type="list",formula1=lst,allow_blank=True); ws.add_data_validation(dv); dv.add(f"{col}2:{col}{N}")
    ws.column_dimensions["S"].hidden=True

# ---- README
rd=wb.active; rd.title="README"
readme=[("VCE General Maths – Exam Question Analysis",None),
("",None),
("What this workbook does",1),
("Every question from each past exam is tagged with a study design topic code and a task type, alongside how well students did on it (from the VCAA examiners' report). The Summary tab turns this into a priority list: which skills cost students the most marks.",None),
("",None),
("Tabs",1),
("Summary – priority list by topic code. Use the yellow filter cells to include only VCAA / NHT / third-party papers, one exam, or a range of years.",None),
("Trends – marks per topic code in each paper, to see year-to-year patterns.",None),
("Master – every question from every paper in one filterable table. Use the column filter arrows (e.g. Source = VCAA NHT, Task type = I).",None),
("Video solution – links to MaffsGuru worked-solution videos. Where the video lists timestamps, the link jumps straight to that question (e.g. ▶ 12:59); otherwise it opens the full video. No videos yet for 2026 NHT or the sample exams. Parts of 2023 Exam 2 jump to the start of the whole question.",None),
("Resources – one row per topic: priority tier, Cambridge sections, two teaching videos, ClassPad and TI-Nspire how-to videos, and every past exam question on that topic to re-attempt (hardest first). This is the relearning lookup for step 3.",None),
("Textbook Map – every Cambridge skills checklist item, mapped to a topic code, with the worked example, the homework questions to redo and the Exam 1 style questions for that section. Filter the Code column to see everything for one topic.",None),
("One tab per paper (e.g. '2025 VCAA E1') – the tagging for that paper. The last column shows which Cambridge sections to redo for each question.",None),
("Codes – the 35 topic codes, taken from the VCE Mathematics Study Design 2023–2027 (General Maths Units 3 & 4, Outcome 1). RF9 (interest-only loans) and MA7 (Leslie-type matrices) are sub-types VCAA examines that the study design does not name separately: RF9 sits under reducing balance loans/recurrence relations, MA7 under transition matrices for populations.",None),
("",None),
("Key measures",1),
("Success rate – Exam 1: % of students choosing the correct answer. Exam 2: average mark ÷ marks available (from the report's score distribution).",None),
("Difficulty – Easy ≥ 70%, Medium 40–69%, Hard < 40%.",None),
("Marks lost – marks × (1 − success rate): the marks the average student dropped on that question.",None),
("Share of exam marks – the average share of an exam's marks that a topic carries, across every included paper (VCAA main, NHT and sample). This is 'how often it comes up'.",None),
("Success rate (Summary) – pooled across every question on that topic from papers with published statistics (the 2023–2025 VCAA main exams; NHT and sample papers have none). 'Marks behind success rate' shows how much data it rests on – red when under 5 marks. A topic with no data uses the average success rate.",None),
("Exam % lost per sitting – share of exam marks × (1 − success rate): roughly the % of the exam score the average student loses on that topic in one sitting. Exam 1 and Exam 2 count equally (each is 30% of the study score). This drives the ranking.",None),
("Priority tiers – ranks 1–9 = Must master, 10–20 = Should master, the rest = Know it. A code with no questions shows 'Not yet examined' – it is still examinable.",None),
("Note: Exam 1 success rates include lucky guesses (a 25% floor), so Exam 1 questions look slightly easier than they really are.",None),
("",None),
("Task types",1),
("C = Calculate   I = Interpret (explain in context)   K = Construct (build a matrix, network, recurrence, table or display)   T = Technology (key step is a CAS / finance solver procedure)",None),
("",None),
("Adding a new exam (e.g. a third-party paper)",1),
("1. Right-click an existing paper tab → Move or Copy → tick 'Create a copy'. Rename it (e.g. '2025 Insight E1').",None),
("2. Overwrite the yellow cells: Paper ID, Year, Source, Exam, Question, Marks, Topic code, Task type, Success rate (leave blank if you have no data). Delete or add rows as needed. Grey columns calculate themselves.",None),
("3. Copy the finished rows (columns A–S) and paste them into the first empty row under the Master table. The Summary and Trends tabs update automatically.",None),
("4. On Trends, copy the last paper column to the right and change the Paper ID in its header cell.",None),
("",None),
("Papers included: VCAA main exams 2023, 2024, 2025; NHT exams 2024, 2025, 2026; VCAA sample exams (current study design). 14 papers, 700 marks. 2023 Exam 2 Q9d was invalidated by VCAA and is excluded from success rates.",None)]
rd.column_dimensions["A"].width=130
for i,(t,h) in enumerate(readme,1):
    c=rd.cell(i,1,t); c.alignment=Alignment(wrap_text=True,vertical="top")
    c.font=Font(name=F,size=16,bold=True,color=NAVY) if i==1 else (Font(name=F,size=12,bold=True,color=NAVY) if h else base)

# ---- Lists
ls=wb.create_sheet("Lists")
for c,(h,vals) in enumerate([("Task type",["C","I","K","T"]),("Source",["VCAA Main","VCAA NHT","VCAA Sample","Third-party"]),("Exam",["Exam 1","Exam 2"]),("Source filter",["All","VCAA Main","VCAA NHT","VCAA Sample","Third-party"]),("Exam filter",["All","Exam 1","Exam 2"])],1):
    ls.cell(1,c,h).font=Font(name=F,bold=True)
    for r,v in enumerate(vals,2): ls.cell(r,c,v).font=base
ls.sheet_state="hidden"

# ---- Codes
cd=wb.create_sheet("Codes")
import json
SK=json.load(open('/tmp/w/skills.json'))
def secs(code):
    seen=[]
    for o in SK:
        if o['code']==code and o['section'] not in seen: seen.append(o['section'])
    return ", ".join(seen)
def rev(code):
    a={"DA1":"6A","DA2":"6A","DA3":"6A","DA4":"6A","DA5":"6B","DA6":"6B","DA7":"6C","DA8":"6C","DA9":"6D","DA10":"6D","DA11":"6D"}
    if code in a: return a[code]+", 6E (Exam 2 style), 16A–16B"
    return {"RF":"9A, 9B (Exam 2 style), 16A–16B","MA":"12A, 12B (Exam 2 style), 16A–16B","ND":"15A, 15B (Exam 2 style), 16A–16B"}[code[:2]]
for c,(h,w) in enumerate([("Code",8),("Unit",8),("Area of study",28),("Topic (study design key knowledge / skills)",80),("Cambridge sections",18),("Cambridge revision sets (exam-style)",30)],1):
    x=cd.cell(1,c,h); x.font=hdr_font; x.fill=hdr_fill; cd.column_dimensions[L(c)].width=w
for r,row in enumerate(CODES,2):
    for c,v in enumerate(list(row)+[secs(row[0]),rev(row[0])],1): cd.cell(r,c,v).font=base
cd.cell(len(CODES)+3,1,"Source: VCAA, VCE Mathematics Study Design 2023–2027, General Mathematics Units 3 and 4, Outcome 1 key knowledge and key skills.").font=Font(name=F,italic=True,size=9)
cd.freeze_panes="A2"

# ---- Textbook map
tm=wb.create_sheet("Textbook Map")
TH=[("Code",8),("Topic",40),("Section",8),("Section title",34),("Checklist skill",60),("Worked example",14),("Homework to redo",18),("Exam 1 style questions",18)]
for c,(h,w) in enumerate(TH,1):
    x=tm.cell(1,c,h); x.font=hdr_font; x.fill=hdr_fill; x.alignment=Alignment(wrap_text=True,vertical="center"); tm.column_dimensions[L(c)].width=w
order={c[0]:i for i,c in enumerate(CODES)}
SKs=sorted(SK,key=lambda o:(order[o['code']],o['chapter'],o['section'],o['n']))
for r,o in enumerate(SKs,2):
    vals=[o['code'],f'=IFERROR(INDEX(Codes!$D:$D,MATCH(A{r},Codes!$A:$A,0)),"")',o['section'],o['section_title'],o['skill'].replace('I can ','',1).capitalize(),o['example'],o['homework'],o['exam_style']]
    for c,v in enumerate(vals,1):
        x=tm.cell(r,c,v); x.font=base; x.alignment=Alignment(wrap_text=True,vertical="top")
tl=len(SKs)+1
t=Table(displayName="TextbookMap",ref=f"A1:H{tl}"); t.tableStyleInfo=TableStyleInfo(name="TableStyleLight9",showRowStripes=True); tm.add_table(t)
tm.freeze_panes="C2"
tm.cell(tl+2,1,"Source: Jones et al., Cambridge Senior Maths VCE General Mathematics Units 3&4 (2023), chapter Skills checklists. Homework = the exercise questions grouped under the same worked example (extracted automatically from the textbook layout; spot-check before giving to students). Textbook typos corrected: 8I checklist points to 'Exercise 8H' (should be 8I); Ch 13 checklist labels trees as 13E (should be 13F).").font=Font(name=F,italic=True,size=9)

# ---- Paper tabs + master
from rows import PAPERS
papers=[(pid,) for pid,_ in PAPERS]
allrows=[]
for pid,rows in PAPERS:
    allrows+=rows
    data_sheet(wb.create_sheet(pid),rows,"T"+pid.replace(" ","_"))
ms=wb.create_sheet("Master",1); data_sheet(ms,allrows,"Master")

# ---- Summary
sm=wb.create_sheet("Summary",1)
sm["A1"]="Priority list by topic code"; sm["A1"].font=Font(name=F,size=16,bold=True,color=NAVY)
lab=[("Source",'All'),("Exam",'All'),("Year from",2022),("Year to",2030)]
for i,(l,v) in enumerate(lab,3):
    sm.cell(i,1,l).font=Font(name=F,bold=True); c=sm.cell(i,2,v); c.fill=inp_fill; c.font=base
dv=DataValidation(type="list",formula1="=Lists!$D$2:$D$6"); sm.add_data_validation(dv); dv.add("B3")
dv=DataValidation(type="list",formula1="=Lists!$E$2:$E$4"); sm.add_data_validation(dv); dv.add("B4")
M=lambda col:f"Master!${col}$2:${col}${N}"
F_=f'{M("C")},IF($B$3="All","*",$B$3),{M("B")},">="&$B$5,{M("B")},"<="&$B$6'
ex=f'{M("D")},IF($B$4="All","*",$B$4)'
info=[("D3","Exam 1 papers included",f'=IF(OR($B$4="All",$B$4="Exam 1"),COUNTIFS({M("S")},1,{M("D")},"Exam 1",{F_}),0)'),
      ("D4","Exam 2 papers included",f'=IF(OR($B$4="All",$B$4="Exam 2"),COUNTIFS({M("S")},1,{M("D")},"Exam 2",{F_}),0)'),
      ("D5","Marks with success data",f'=SUMIFS({M("F")},{M("M")},"<>",{ex},{F_})'),
      ("D6","Average success rate (all topics)",f'=IFERROR((D5*0+SUMIFS({M("F")},{M("M")},"<>",{ex},{F_})-SUMIFS({M("Q")},{ex},{F_}))/SUMIFS({M("F")},{M("M")},"<>",{ex},{F_}),0.6)')]
for addr,label,f in info:
    r=int(addr[1:]); sm.cell(r,4,label).font=base; x=sm.cell(r,7,f); x.font=Font(name=F,bold=True)
sm["G6"].number_format="0%"
sm["A7"]="Yellow cells are filters. How often a topic appears uses every included paper; success rates come only from papers with published statistics (VCAA main exams). See README."; sm["A7"].font=Font(name=F,italic=True,size=9)
H=["Code","Area of study","Topic","Questions","Marks","Questions per sitting","Share of exam marks","Success rate","Marks behind success rate","Exam % lost per sitting","Rank","Priority tier","Calculate","Interpret","Construct","Technology"]
W=[8,24,52,10,8,10,10,9,10,11,7,24,9,9,9,10]
HR=9
for c,h in enumerate(H,1):
    x=sm.cell(HR,c,h); x.font=hdr_font; x.fill=hdr_fill; x.alignment=Alignment(wrap_text=True,vertical="center"); sm.column_dimensions[L(c)].width=W[c-1]
sm.row_dimensions[HR].height=45
first=HR+1; last=HR+len(CODES)
for i,(code,u,aos,topic) in enumerate(CODES):
    r=first+i; crit=f'{M("G")},$A{r},{ex},{F_}'
    sh=lambda e,n:f'IF(${n}>0,SUMIFS({M("F")},{M("G")},$A{r},{M("D")},"{e}",{F_})/SUMIFS({M("F")},{M("D")},"{e}",{F_}),0)'
    vals=[code,aos,topic,
      f'=COUNTIFS({crit})',
      f'=SUMIFS({M("F")},{crit})',
      f'=IFERROR(COUNTIFS({M("G")},$A{r},{M("D")},"Exam 1",{F_})/$G$3,0)*($G$3>0)+IFERROR(COUNTIFS({M("G")},$A{r},{M("D")},"Exam 2",{F_})/$G$4,0)*($G$4>0)',
      f'=IFERROR(({sh("Exam 1","G$3")}+{sh("Exam 2","G$4")})/(($G$3>0)+($G$4>0)),0)',
      f'=IF(I{r}>0,(I{r}-SUMIFS({M("Q")},{crit}))/I{r},$G$6)',
      f'=SUMIFS({M("F")},{M("M")},"<>",{crit})',
      f'=G{r}*(1-H{r})',
      f'=IF(D{r}=0,"",RANK(J{r},$J${first}:$J${last}))',
      f'=IF(D{r}=0,"Not yet examined (still examinable)",IF(K{r}<=9,"Must master",IF(K{r}<=20,"Should master","Know it")))']
    for t in ["C","I","K","T"]: vals.append(f'=COUNTIFS({crit},{M("K")},"{t}")')
    for c,v in enumerate(vals,1):
        x=sm.cell(r,c,v); x.font=base; x.border=Border(bottom=thin)
    sm.cell(r,6).number_format="0.0"; sm.cell(r,7).number_format="0.0%"; sm.cell(r,8).number_format="0%"; sm.cell(r,10).number_format="0.00%"
    sm.cell(r,8).comment=None
rng=f"L{first}:L{last}"
for txt,col in (("Must master","F8CBAD"),("Should master","FFE699"),("Know it","C6EFCE"),("Not yet","D9D9D9")):
    sm.conditional_formatting.add(rng,FormulaRule(formula=[f'LEFT(L{first},{len(txt)})="{txt}"'],fill=PatternFill("solid",fgColor=col)))
sm.conditional_formatting.add(f"H{first}:H{last}",ColorScaleRule(start_type="num",start_value=0.3,start_color="F8696B",mid_type="num",mid_value=0.55,mid_color="FFEB84",end_type="num",end_value=0.85,end_color="63BE7B"))
sm.conditional_formatting.add(f"J{first}:J{last}",ColorScaleRule(start_type="min",start_color="FFFFFF",end_type="max",end_color="F8696B"))
sm.conditional_formatting.add(f"I{first}:I{last}",CellIsRule(operator="lessThan",formula=["5"],font=Font(name=F,color="C00000",italic=True)))
ar=last+3
sm.cell(ar-1,1,"By area of study").font=Font(name=F,size=12,bold=True,color=NAVY)
for c,h in enumerate(["","Area of study","","Questions","Marks","Questions per sitting","Share of exam marks","Success rate","Marks behind success rate","Exam % lost per sitting"],1):
    if h: x=sm.cell(ar,c,h); x.font=hdr_font; x.fill=hdr_fill; x.alignment=Alignment(wrap_text=True)
for i,a in enumerate(["Data analysis","Recursion & financial modelling","Matrices","Networks & decision maths"]):
    r=ar+1+i; B_=lambda col:f'SUMIF($B${first}:$B${last},$B{r},{col}${first}:{col}${last})'
    sm.cell(r,2,a)
    for c,col in ((4,"$D"),(5,"$E"),(6,"$F"),(7,"$G"),(9,"$I"),(10,"$J")): sm.cell(r,c,f'={B_(col)}')
    sm.cell(r,8,f'=IFERROR(SUMPRODUCT(($B${first}:$B${last}=$B{r})*($I${first}:$I${last})*($H${first}:$H${last}))/I{r},"")')
    for c in range(2,11): sm.cell(r,c).font=base
    sm.cell(r,6).number_format="0.0"; sm.cell(r,7).number_format="0.0%"; sm.cell(r,8).number_format="0%"; sm.cell(r,10).number_format="0.0%"
sm.freeze_panes=f"D{HR+1}"
sm.auto_filter.ref=f"A{HR}:P{last}"

# ---- Resources
from res import V as RV, R as RR, CP, TI
rs=wb.create_sheet("Resources")
RH=[("Code",7),("Topic",38),("Priority tier",16),("Cambridge sections",14),("Cambridge revision (exam-style)",20),("Teaching video 1",34),("Teaching video 2",34),("ClassPad how-to",28),("TI-Nspire how-to",28),("Past exam questions to re-attempt (hardest first)",60),("Timestamped walkthrough clips",12)]
for c,(h,w) in enumerate(RH,1):
    x=rs.cell(1,c,h); x.font=hdr_font; x.fill=hdr_fill; x.alignment=Alignment(wrap_text=True,vertical="center"); rs.column_dimensions[L(c)].width=w
rs.row_dimensions[1].height=32
def short(pid): return pid.replace(" VCAA","").replace("VCAA Sample","Sample")
def vcell(r,c,vid):
    if not vid: return
    t,ch,ln=RV[vid]; x=rs.cell(r,c,f"{t} — {ch} ({ln})"); x.hyperlink=f"https://www.youtube.com/watch?v={vid}"
    x.font=Font(name=F,size=10,color="0563C1",underline="single")
for i,(code,u,aos,topic) in enumerate(CODES):
    r=i+2
    qs=[(row[12] if row[12] is not None else 2, f"{short(row[0])} Q{row[4]}"+(f" ({row[12]:.0%})" if row[12] is not None else "")) for pid,rows in PAPERS for row in rows if row[6]==code]
    qs.sort(key=lambda x:x[0])
    vals=[code,f'=IFERROR(INDEX(Codes!$D:$D,MATCH(A{r},Codes!$A:$A,0)),"")',f'=IFERROR(INDEX(Summary!$L:$L,MATCH(A{r},Summary!$A:$A,0)),"")',
          f'=IFERROR(INDEX(Codes!$E:$E,MATCH(A{r},Codes!$A:$A,0)),"")',f'=IFERROR(INDEX(Codes!$F:$F,MATCH(A{r},Codes!$A:$A,0)),"")',None,None,None,None,
          "; ".join(q for _,q in qs),f'=COUNTIFS(Master!$G$2:$G${N},A{r},Master!$U$2:$U${N},"▶*")']
    for c,v in enumerate(vals,1):
        x=rs.cell(r,c,v); x.font=base; x.alignment=Alignment(wrap_text=True,vertical="top")
    vcell(r,6,RR[code][0]); vcell(r,7,RR[code][1] if len(RR[code])>1 else None); vcell(r,8,CP.get(code)); vcell(r,9,TI.get(code))
    for c in (6,7,8,9): rs.cell(r,c).alignment=Alignment(wrap_text=True,vertical="top")
    if not CP.get(code): rs.cell(r,8,"Not needed – by hand").font=Font(name=F,size=9,italic=True,color="808080")
    if not TI.get(code): rs.cell(r,9,"Not needed – by hand").font=Font(name=F,size=9,italic=True,color="808080")
    if code=="ND7": rs.cell(r,9,"TI-Nspire Hungarian algorithm program: see MaffsGuru 2023 Exam 1 video description").font=Font(name=F,size=9,italic=True)
lr=len(CODES)+1
for txt,col in (("Must master","F8CBAD"),("Should master","FFE699"),("Know it","C6EFCE"),("Not yet","D9D9D9")):
    rs.conditional_formatting.add(f"C2:C{lr}",FormulaRule(formula=[f'LEFT(C2,{len(txt)})="{txt}"'],fill=PatternFill("solid",fgColor=col)))
rs.freeze_panes="C2"; rs.auto_filter.ref=f"A1:K{lr}"
rs.cell(lr+2,1,"Teaching videos chosen from YouTube searches (Sept 2026): current-study-design VCE content preferred, full videos only (MaffsGuru 'Preview' clips excluded). Re-attempt list: every past question on that topic, hardest first; % = cohort success rate (none for NHT/sample). Walkthrough clips: filter the Master tab by Code and click the Video solution column.").font=Font(name=F,italic=True,size=9)

# ---- Trends
tr=wb.create_sheet("Trends",2)
tr["A1"]="Marks per topic code in each paper"; tr["A1"].font=Font(name=F,size=16,bold=True,color=NAVY)
tr["A2"]="Each column header is a Paper ID from the Master tab. To add a paper, copy the last column to the right and change its header."; tr["A2"].font=Font(name=F,italic=True,size=9)
tr.cell(4,1,"Code"); tr.cell(4,2,"Topic")
for c in (1,2):
    x=tr.cell(4,c); x.font=hdr_font; x.fill=hdr_fill
tr.column_dimensions["A"].width=8; tr.column_dimensions["B"].width=60
for j,(pid,*_) in enumerate(papers):
    c=3+j; x=tr.cell(4,c,pid); x.font=hdr_font; x.fill=hdr_fill; x.alignment=Alignment(wrap_text=True); tr.column_dimensions[L(c)].width=13
    for i,(code,*_x) in enumerate(CODES):
        r=5+i; tr.cell(r,c,f'=SUMIFS({M("F")},{M("G")},$A{r},{M("A")},{L(c)}$4)').font=base
for i,(code,u,aos,topic) in enumerate(CODES):
    tr.cell(5+i,1,code).font=base; tr.cell(5+i,2,topic).font=base
tl=5+len(CODES)
tr.conditional_formatting.add(f"C5:Z{tl-1}",ColorScaleRule(start_type="num",start_value=0,start_color="FFFFFF",end_type="num",end_value=6,end_color="5B9BD5"))
tr.freeze_panes="C5"

for t,r in LINKS: wb[t].cell(r,21).font=Font(name=F,size=10,color='0563C1',underline='single')
for ws in wb.worksheets: ws.sheet_view.zoomScale=100
order=["README","Summary","Resources","Trends","Master","Textbook Map"]+[p for p,_ in PAPERS]+["Codes","Lists"]
wb._sheets=[wb[n] for n in order]; wb.active=0
wb.save("/tmp/w/GM_Exam_Analysis.xlsx")
