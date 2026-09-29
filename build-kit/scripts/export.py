import sys,json,re; sys.path.insert(0,'/tmp/w')
from rows import PAPERS
from data import CODES
from res import V,R,CP,TI
from video import link as VL
from openpyxl import load_workbook
R_=json.load(open('reports.json')); SK=json.load(open('skills.json'))
norm=lambda q:re.sub(r'[^0-9a-z]','',q.lower())
RK={"2023 VCAA E1":"2023 - Exam 1","2023 VCAA E2":"2023 - Exam 2","2024 VCAA E1":"2024 - Exam 1","2024 VCAA E2":"2024 - Exam 2","2025 VCAA E1":"2025 - Exam 1","2025 VCAA E2":"2025 - Exam 2","2024 NHT E2":"2024 NHT - Exam 2","2025 NHT E2":"2025 NHT - Exam 2","2026 NHT E2":"2026 NHT - Exam 2"}
from docx import Document
B="/mnt/user-data/outputs/org/02 Past Exams/"
def nhtE1(folder):
    d=Document(B+f"{folder}/{folder} - Exam 1 Report.docx");a={}
    for t in d.tables:
        for r in t.rows:
            c=[x.text.strip() for x in r.cells]
            if c[0].isdigit(): a[c[0]]=c[-1]
    return a
NE1={"2024 NHT E1":nhtE1("2024 NHT"),"2025 NHT E1":nhtE1("2025 NHT"),"2026 NHT E1":nhtE1("2026 NHT")}
def clean(s):
    s=re.sub(r'\$[^$]*\$','',s or ''); s=re.sub(r'\[table\][^\]]*?(?= [A-Z][a-z]|$)','',s)
    s=re.sub(r'\s+',' ',s).strip()
    return s[:320]+('…' if len(s)>320 else '')
def comment(pid,q,exam):
    if pid in NE1: return clean(NE1[pid].get(q,''))
    rep=R_.get(RK.get(pid,''),[])
    if exam=="Exam 1":
        row=next((r for r in rep if r[0]==q),None)
        if row and not re.fullmatch(r'[\d.]+',row[-1]): return clean(row[-1])
        return ''
    row=next((r for r in rep if norm(r['q'])==norm(q) and r.get('stats')),None)
    return clean(row['text']) if row else ''
# summary numbers from workbook
wb=load_workbook('GM_Exam_Analysis.xlsx',data_only=True); s=wb['Summary']
SUM={}
for r in s.iter_rows(min_row=10,max_row=44,values_only=True):
    SUM[r[0]]=dict(q=r[3],marks=r[4],per=round(r[5],1),share=round(r[6],4),succ=round(r[7],3),lost=round(r[9],4),rank=r[10],tier=r[11])
qs=[]
for pid,rows in PAPERS:
    for row in rows:
        u,lab=VL(pid,row[4])
        qs.append(dict(p=pid,y=row[1],src=row[2],ex=row[3],q=row[4],m=row[5],c=row[6],c2=row[9],t=row[10],a=row[11],s=row[12],sp=row[14],n=row[15],v=u,vl=lab,cm=comment(pid,row[4],row[3])))
topics={}
for code,u,aos,topic in CODES:
    sk=[dict(sec=o['section'],title=o['section_title'],skill=o['skill'].replace('I can ','',1),ex=o['example'],hw=o['homework'],es=o['exam_style']) for o in SK if o['code']==code]
    rev={"DA1":"6A","DA2":"6A","DA3":"6A","DA4":"6A","DA5":"6B","DA6":"6B","DA7":"6C","DA8":"6C","DA9":"6D","DA10":"6D","DA11":"6D"}.get(code)
    rev=(rev+", 6E (Exam 2 style)") if rev else {"RF":"9A, 9B (Exam 2 style)","MA":"12A, 12B (Exam 2 style)","ND":"15A, 15B (Exam 2 style)"}[code[:2]]
    vid=lambda i:dict(id=i,t=V[i][0],ch=V[i][1],len=V[i][2]) if i else None
    topics[code]=dict(name=topic,aos=aos,unit=u,sum=SUM[code],vids=[vid(i) for i in R[code]],cp=vid(CP.get(code)),ti=vid(TI.get(code)),skills=sk,rev=rev+", 16A–16B")
json.dump(dict(qs=qs,topics=topics),open('appdata.json','w'),ensure_ascii=False,separators=(',',':'))
print(len(qs), sum(1 for q in qs if q['cm']), len(open('appdata.json').read())//1024,'KB')
print([q for q in qs if q['p']=='2025 VCAA E2'][18])
