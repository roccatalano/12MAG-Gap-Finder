import re,json,sys,glob,os
from docx import Document
from docx.oxml.ns import qn
B="/mnt/user-data/outputs/org/02 Past Exams/"
def txt(el): return ''.join(x.text or '' for x in el.iter(qn('w:t'))).strip()
def e1(path):
    d=Document(path);out=[]
    for t in d.tables:
        for r in t.rows:
            c=[x.text.strip().replace('\n',' ') for x in r.cells]
            if c and re.fullmatch(r'\d{1,2}',c[0]) and len(c)>=6:
                out.append(c)
    return out
def e2(path):
    d=Document(path);out=[];cur=None
    for el in d.element.body:
        if el.tag==qn('w:p'):
            t=txt(el)
            m=re.match(r'^Question\s*(\d+[a-z]?\.?(?:\s?[ivx]+\.?)?)',t)
            if m: cur={'q':m.group(1),'stats':None,'text':''}; out.append(cur); continue
            if cur and t and len(cur['text'])<300: cur['text']+=' '+t
        elif el.tag==qn('w:tbl'):
            rows=[[txt(tc) for tc in tr.iter(qn('w:tc'))] for tr in el.iter(qn('w:tr'))]
            if rows and rows[0] and rows[0][0].startswith('Mark') and cur: cur['stats']=rows
            elif cur and len(cur['text'])<300: cur['text']+=' [table] '+' / '.join('|'.join(r) for r in rows[:3])[:150]
    return out
res={}
for f in sorted(glob.glob(B+"*/*Report.docx")):
    k=os.path.basename(f).replace(' Report.docx','')
    res[k]=e1(f) if 'Exam 1' in k else e2(f)
    print(k,len(res[k]))
json.dump(res,open('/tmp/w/reports.json','w'))
